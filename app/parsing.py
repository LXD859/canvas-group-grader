"""Read untrusted spreadsheets without executing formulas or changing originals."""
import csv
import io
import re
import zipfile
from collections import Counter
from decimal import Decimal, InvalidOperation, localcontext

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLS = 200
DEDUCTION = re.compile(r"^[-−–—﹣－]\s*(\d+(?:\.\d+)?|\.\d+)\s*[,，]\s*(\S[\s\S]*)$")


class InputError(ValueError):
    pass


def text(value):
    return "" if value is None else str(value).strip()


def number(value):
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise InputError("满分必须是有效数字。") from None
    if not result.is_finite() or result < 0:
        raise InputError("满分必须是非负有限数字。")
    return result


def read_table(filename, content, sheet=None, encoding="auto"):
    if len(content) > MAX_BYTES:
        raise InputError("文件超过 10 MB。")
    if filename.lower().endswith(".csv"):
        if encoding not in {"auto", "utf-8-sig", "gb18030"}:
            raise InputError("不支持该编码。")
        decoded = None
        for enc in (["utf-8-sig", "gb18030"] if encoding == "auto" else [encoding]):
            try:
                decoded = content.decode(enc)
                break
            except UnicodeDecodeError:
                pass
        if decoded is None:
            raise InputError("无法解码 CSV，请选择 UTF-8 或 GB18030。")
        try:
            rows = []
            for row in csv.reader(io.StringIO(decoded), strict=True):
                rows.append(row)
                if len(rows) > MAX_ROWS + 1 or len(row) > MAX_COLS:
                    raise InputError("文件最多支持 5000 行数据和 200 列。")
        except csv.Error:
            raise InputError("CSV 格式有误，请检查引号和分隔符。") from None
        sheets, selected = ["CSV"], "CSV"
    elif filename.lower().endswith(".xlsx"):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                    raise InputError("Excel 解压后超过 50 MB。")
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheets = book.sheetnames
                selected = sheet or ("Assignment Grading" if "Assignment Grading" in sheets else sheets[0])
                if selected not in sheets:
                    raise InputError("工作表不存在。")
                tab = book[selected]
                # Some exporters declare A1:A1 even when the sheet contains many rows.
                tab.reset_dimensions()
                rows = []
                for row in tab.iter_rows(values_only=True):
                    rows.append(list(row))
                    if len(rows) > MAX_ROWS + 1 or len(row) > MAX_COLS:
                        raise InputError("文件最多支持 5000 行数据和 200 列。")
            finally:
                book.close()
        except InputError:
            raise
        except Exception:
            raise InputError("无法读取 Excel，请使用未加密的 .xlsx 文件。") from None
    else:
        raise InputError("仅支持 .xlsx 和 .csv 文件。")
    if not rows or not any(text(x) for x in rows[0]):
        raise InputError("第一行必须包含列标题。")
    width = max(map(len, rows))
    rows = [[text(x) for x in r] + [""] * (width - len(r)) for r in rows]
    headers = rows[0]
    return {"sheets": sheets, "sheet": selected, "headers": headers, "rows": rows[1:]}


def calculate_rows(table, group_column, question_columns, maximum):
    maximum = number(maximum)
    width = len(table["headers"])
    if not question_columns or len(set(question_columns)) != len(question_columns):
        raise InputError("请选择至少一个不重复的题目列。")
    indices = [group_column, *question_columns]
    if any(i < 0 or i >= width for i in indices) or group_column in question_columns:
        raise InputError("小组列和题目列必须有效且互不重叠。")
    question_columns = sorted(question_columns)
    labels = [table["headers"][i] for i in question_columns]
    if any(not label for label in labels) or len(set(labels)) != len(labels):
        raise InputError("所选题目列必须有唯一的题号标题。")
    names = Counter(r[group_column] for r in table["rows"] if r[group_column])
    results = []
    for row_number, row in enumerate(table["rows"], 2):
        if not any(row):
            continue
        name, errors, items, comments = row[group_column], [], [], []
        if not name:
            errors.append(f"{get_column_letter(group_column + 1)}{row_number}：缺少小组名称。")
        elif names[name] > 1:
            errors.append("表格中存在重复小组名称。")
        if name.startswith("="):
            errors.append("小组名称必须为文本，不能使用公式。")
        with localcontext() as context:
            context.prec = 64
            total = Decimal(0)
            for index in question_columns:
                raw, label = row[index], table["headers"][index]
                coordinate = f"{get_column_letter(index + 1)}{row_number}"
                if not raw:
                    items.append({"question": label, "deduction": "0", "raw": "", "cell": coordinate})
                    continue
                match = DEDUCTION.fullmatch(raw)
                if not match or len(match.group(1)) > 24:
                    errors.append(f"{coordinate}（{label}）：请使用 -1.5, English reason 格式。")
                    continue
                deduction = Decimal(match.group(1))
                total += deduction
                items.append({"question": label, "deduction": str(deduction), "raw": raw, "cell": coordinate})
                comments.append(f"{label}: {raw}")
            score = maximum - total
        if score < 0:
            errors.append("累计扣分超过作业满分。")
        comment = "\n".join(comments)
        if len(comment) > 16000:
            errors.append("合并评语超过 16000 字符。")
        results.append({"row": row_number, "name": name, "items": items,
                        "deduction": str(total), "score": str(score), "comment": comment,
                        "errors": errors})
    if not results:
        raise InputError("工作表中没有小组数据。")
    return results
